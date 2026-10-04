"""Public regressions from the October 2026 independent review."""

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from digital_metrics import Evaluation, Metrics
from digital_metrics.ci import calculate_confidence_interval
from digital_metrics.matching import match_boxes
from digital_metrics.reporting import get_dashboards, plot_confidence_intervals
from digital_metrics.scoring import compute_ap, compute_kappa, get_confusion_matrix
from digital_metrics.translit import restore_labels


def frames() -> tuple[pd.DataFrame, pd.DataFrame]:
    gt = pd.DataFrame(
        [["img", "cat", 0.0, 0.0, 10.0, 10.0, "test"]],
        columns=[
            "image_name",
            "instance_label",
            "bbox_x_tl",
            "bbox_y_tl",
            "bbox_x_br",
            "bbox_y_br",
            "split",
        ],
    )
    preds = gt.drop(columns="split").assign(confidence=0.9)
    return gt, preds


def test_mutable_ci() -> None:
    m = Metrics(tp=1)
    _ = m.precision_ci_lower, m.recall_ci_lower
    m.fp = m.fn = 99
    assert m.model_dump()["precision_ci_upper"] < 0.1
    assert m.recall_ci_upper < 0.1


@pytest.mark.parametrize("level", [-0.1, 0.0, 1.0, 1.1, float("nan")])
def test_ci_rejects_invalid_levels_even_empty(level: float) -> None:
    with pytest.raises(ValueError):
        calculate_confidence_interval(0, 0, confidence_level=level)


def test_ci_rejects_invalid_method_even_empty() -> None:
    with pytest.raises(ValueError):
        calculate_confidence_interval(0, 0, method="bootstrap")


def test_plot_coverage(tmp_path: Path) -> None:
    m = {"cat": Metrics(tp=7, fp=3)}
    _, a = plot_confidence_intervals(m, "precision", 0.95, str(tmp_path / "95.png"))
    _, b = plot_confidence_intervals(m, "precision", 0.99, str(tmp_path / "99.png"))
    assert a is not None and b is not None
    assert not np.array_equal(a.collections[0].get_segments(), b.collections[0].get_segments())


def test_empty_dashboard_schema(tmp_path: Path) -> None:
    gt, _ = frames()
    devs, dtrk = get_dashboards({}, gt.iloc[:0], None, [], path=str(tmp_path))
    assert devs.empty and dtrk.empty
    assert "confidence" in devs and "Порог" in dtrk
    assert not list(tmp_path.glob("*.png"))


def test_excel_literal_labels_and_suffix(tmp_path: Path) -> None:
    gt, _ = frames()
    gt["instance_label"] = "=1+1"
    for suffix in ["a", "b"]:
        get_dashboards(
            {"=1+1": Metrics(tp=1)},
            gt,
            np.eye(2, dtype=np.int64),
            ["=1+1", "background"],
            suffix,
            path=str(tmp_path),
        )
        for name in ["full_dashboard", "метрики_дтрк", "matrix"]:
            ws = load_workbook(tmp_path / f"{name}_{suffix}.xlsx").active
            assert ws is not None and ws["A2"].value == "=1+1"
            assert ws["A2"].data_type == "s"
        ws = load_workbook(tmp_path / f"matrix_{suffix}.xlsx").active
        assert ws is not None and ws["B1"].data_type == "s"
    assert len(list(tmp_path.glob("*.png"))) == 8


@pytest.mark.parametrize("strategy", ["greedy", "iou_prior", "hungarian"])
def test_wrong_class_one_cm_event(strategy: str) -> None:
    gt, preds = frames()
    preds["instance_label"] = "dog"
    matches = match_boxes(gt, preds, 0.5, strategy=strategy)
    cm, _ = get_confusion_matrix(matches, ["cat", "dog"])
    assert cm[0, 1] == 1 and cm.sum() == 1
    preds = pd.concat([preds, preds], ignore_index=True)
    cm, _ = get_confusion_matrix(match_boxes(gt, preds, 0.5, strategy=strategy), ["cat", "dog"])
    assert cm[0, 1] == 1 and cm[2, 1] == 1 and cm.sum() == 2


def test_numeric_ids_and_collision() -> None:
    gt, preds = frames()
    gt["image_name"] = preds["image_name"] = 7
    assert match_boxes(gt, preds, 0.5)["cat"][0].type == "TP"
    gt = pd.concat([gt, gt.assign(image_name="7")], ignore_index=True)
    with pytest.raises(ValueError, match="collision"):
        Evaluation(preds, gt)


@pytest.mark.parametrize("index", ["box-uuid", pd.Timestamp("2026-10-04"), ("part", 2)])
def test_preserve_index(index: object) -> None:
    gt, preds = frames()
    gt.index = pd.Index([index], tupleize_cols=False)
    preds.index = pd.Index([index], tupleize_cols=False)
    match = match_boxes(gt, preds, 0.5)["cat"][0]
    assert match.pred_index == index and match.gt_index == index


@pytest.mark.parametrize("bad", [float("nan"), float("inf")])
def test_raw_validation_before_preprocessing(bad: float) -> None:
    gt, preds = frames()
    preds["confidence"] = bad
    with pytest.raises(ValueError, match="confidence"):
        Evaluation(preds, gt, preprocess_preds_conf_threshold=0.1)()
    gt.loc[0, "bbox_x_tl"] = bad
    with pytest.raises(ValueError, match="Ground-truth"):
        Evaluation(frames()[1], gt)()


def test_background_reserved() -> None:
    gt, preds = frames()
    gt["instance_label"] = "background"
    with pytest.raises(ValueError, match="background"):
        match_boxes(gt, preds, 0.5)


def test_kappa_clipping() -> None:
    assert compute_kappa([[-2, 0, 2, 2]], [[-2, 0, 2, 2]], (10, 10)) == 1


def test_empty_evaluation_and_split_validation() -> None:
    gt, preds = frames()
    ev = Evaluation(preds.iloc[:0], gt.iloc[:0])
    ev(find_best_confs=False)
    assert ev.metrics == {} and ev.cm is not None and ev.cm.tolist() == [[0]]
    with pytest.raises(ValueError, match="split"):
        Evaluation(preds, gt)("missing")
    with pytest.raises(ValueError, match="split"):
        Evaluation(preds, gt.drop(columns="split"))("test")
    with pytest.raises(ValueError, match="split"):
        Evaluation(preds, pd.concat([gt, gt.assign(split="val")], ignore_index=True))


@pytest.mark.parametrize(
    "kw",
    [
        dict(matching_strategy="oops"),
        dict(ap_method="oops"),
        dict(confidence_optimization="oops"),
        dict(iou_threshold=-0.1),
        dict(iou_threshold=float("nan")),
    ],
)
def test_config_validation(kw: dict[str, object]) -> None:
    gt, preds = frames()
    with pytest.raises(ValueError):
        Evaluation(preds, gt, **kw)


def test_public_ap_validation() -> None:
    with pytest.raises(ValueError):
        compute_ap(np.array([1.0]), np.array([1.0]), method="typo")


def test_split_visualization_and_dedup() -> None:
    gt, preds = frames()
    other = gt.assign(image_name="other", split="val")
    ev = Evaluation(
        pd.concat([preds, preds.assign(image_name="other")], ignore_index=True),
        pd.concat([gt, other], ignore_index=True),
    )
    ev("test", find_best_confs=False)
    assert set(ev.get_dfs_visualization()[1]["image_name"]) == {"img"}
    gt = pd.concat([gt, gt.assign(instance_label="dog")], ignore_index=True)
    ev = Evaluation(preds, gt, preprocess=True)
    assert len(ev.split_df) == 2


def test_ambiguous_translit_stays_unresolved() -> None:
    assert restore_labels(["shch"], ["щ", "шч"])["shch"] == "shch"


def test_missing_label_column() -> None:
    gt, preds = frames()
    with pytest.raises(ValueError, match="instance_label"):
        Evaluation(preds.drop(columns="instance_label"), gt)()


@pytest.mark.parametrize("index", [-1, None, pd.NA, pd.NaT])
def test_reject_missing_or_sentinel_row_label(index: object) -> None:
    gt, preds = frames()
    gt.index = pd.Index([index], tupleize_cols=False)
    with pytest.raises(ValueError, match="row labels"):
        match_boxes(gt, preds, 0.5)


def test_label_preservation_through_filter_and_visualization() -> None:
    gt, preds = frames()
    gt.index = pd.Index([("gt", 1)], tupleize_cols=False)
    preds.index = pd.Index(["known-box"])
    unknown = preds.assign(instance_label="unknown")
    unknown.index = pd.Index(["unknown-box"])
    ev = Evaluation(
        pd.concat([preds, unknown]), gt, preprocess=True, preprocess_preds_conf_threshold=0.2
    )
    ev(find_best_confs=False)
    assert ev.metrics["cat"].tp == 1
    g, p = ev.get_dfs_visualization()
    assert g.index.tolist() == [("gt", 1)] and p.index.tolist() == ["known-box"]
    assert p.predict_type.tolist() == ["TP"]


def test_exact_zero_tp_backend_adapter() -> None:
    from digital_metrics import DetectionMetrics
    from digital_metrics.calibration import ConfidenceCalibrator
    from digital_metrics.engines.backend import BackendEngine

    gt, _ = frames()
    engine = BackendEngine(
        backend="torchmetrics",
        classes=["cat"],
        confidence_optimization="per_class",
        calibrator=ConfidenceCalibrator(
            classes=["cat"],
            iou_threshold=0.5,
            matching_strategy="greedy",
            confidence_optimization="per_class",
        ),
    )
    value = DetectionMetrics(
        precision=0, recall=0, f1=0, ap50=0, ap75=0, ap50_95=0, tp=0, fp=4, fn=1
    )
    metric = engine._adapt({"cat": value}, gt)["cat"]
    assert metric.counts_observed and metric.fp == 4 and metric.fn == 1
    assert metric.precision_ci_upper > 0


def test_warning_filters_are_not_changed_by_import() -> None:
    import subprocess
    import sys

    subprocess.run(
        [
            sys.executable,
            "-c",
            "import warnings; warnings.simplefilter('always', RuntimeWarning); "
            "import digital_metrics; "
            "warnings.warn('external-warning-visible', RuntimeWarning)",
        ],
        check=True,
        capture_output=True,
        text=True,
    ).stderr.index("external-warning-visible")


@pytest.mark.parametrize("failure", [False, True])
def test_inference_releases_results_before_allocator_cleanup(
    monkeypatch: pytest.MonkeyPatch, *, failure: bool
) -> None:
    import sys
    import types
    import weakref

    from digital_metrics.inference import yolo_predict

    refs: list[weakref.ReferenceType[object]] = []

    class Result:
        boxes = None

    def make_result() -> Result:
        result = Result()
        refs.append(weakref.ref(result))
        return result

    class Model:
        names: dict[int, str] = {}

        def predict(self, **_: object) -> object:
            yield make_result()
            if failure:
                raise RuntimeError("original inference error")

    monkeypatch.setitem(sys.modules, "ultralytics", types.SimpleNamespace(YOLO=lambda _: Model()))

    def cleanup(_: object) -> None:
        assert refs and all(ref() is None for ref in refs)

    monkeypatch.setattr(yolo_predict, "_release_gpu_memory", cleanup)
    if failure:
        with pytest.raises(RuntimeError, match="original inference error"):
            yolo_predict.predict_on_images("weights", ["image"])
    else:
        assert yolo_predict.predict_on_images("weights", ["image"]).empty


def test_failed_tracker_attempts_independent_cleanup(monkeypatch: pytest.MonkeyPatch) -> None:
    from digital_metrics import ClearMLTracker

    events: list[str] = []

    class Task:
        status = "running"

        def get_logger(self) -> object:
            return object()

        def mark_failed(self, **_: object) -> None:
            events.append("failed")
            raise RuntimeError("status service error")

        def close(self) -> None:
            events.append("close")

    tracker = ClearMLTracker(task=Task(), attach_logs=False)

    def detach() -> None:
        events.append("detach")
        raise RuntimeError("sink error")

    monkeypatch.setattr(tracker, "detach_loguru", detach)
    with pytest.raises(ValueError, match="body error"):
        with tracker:
            raise ValueError("body error")
    assert events == ["detach", "close", "failed"]


@pytest.mark.parametrize("right", [0.0, -1.0])
def test_labelled_boxes_require_positive_area(right: float) -> None:
    gt, preds = frames()
    gt["bbox_x_br"] = right
    with pytest.raises(ValueError, match="positive area"):
        Evaluation(preds, gt)


def test_failed_tracker_does_not_mark_failed_before_sdk_close() -> None:
    from digital_metrics import ClearMLTracker

    events: list[str] = []

    class Task:
        def get_logger(self) -> object:
            return object()

        def close(self) -> None:
            events.append("close")
            raise RuntimeError("watchdog shutdown failed")

        def mark_failed(self, **_: object) -> None:
            events.append("unsafe failure transition")

    tracker = ClearMLTracker(task=Task(), attach_logs=False)
    original = ValueError("body failure")
    with pytest.raises(ValueError) as raised, tracker:
        raise original
    assert raised.value is original
    assert events == ["close"]
